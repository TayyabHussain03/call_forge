"""Tests for Slice 42 — Deterministic Excel/CSV Result Writer.

Proves: XLSX/CSV export preserves original data, maps outcomes via stable
identity, serializes collections deterministically, is idempotent on rerun,
never corrupts the source file, keeps DNC distinct, and separates
customer outcome from execution status.
"""

from __future__ import annotations

import csv
import os
import tempfile

import openpyxl
import pytest

from app.export.contracts import (
    COLLECTION_DELIMITER,
    ExportStatus,
    LeadOutcomeExportResult,
    MultipleOutcomePolicy,
    OutputFormat,
    RESULT_COLUMNS,
)
from app.export.writer import LeadOutcomeExportWriter
from app.leads.contracts import (
    LeadRecord,
    NormalizedPhone,
    PhoneNormalizationStatus,
    WebsiteStatus,
)
from app.runtime.chat.outcome.contracts import (
    DecisionMakerStatus,
    DNCStatus,
    ExecutionStatus,
    FieldConfidence,
    FollowUpNeed,
    IdentifiedProblem,
    InterestLevel,
    LeadOutcome,
    OutcomeStatus,
    OutcomeSummaryLine,
    RecordedObjection,
    RelevantService,
    TerminationReason,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _phone(raw: str = "+15551234567") -> NormalizedPhone:
    return NormalizedPhone(
        raw=raw, normalized=raw, status=PhoneNormalizationStatus.VALID_E164,
    )


def _lead(
    lead_id: str = "lead-1",
    row_number: int = 1,
    phone: str = "+15551234567",
    business_name: str = "Alpha Inc",
) -> LeadRecord:
    return LeadRecord(
        lead_id=lead_id,
        tenant_id="t1",
        campaign_id="c1",
        business_name=business_name,
        phone=_phone(phone),
        category=None,
        city=None, state=None, country=None, address=None,
        email=None, contact_name=None,
        import_batch_id="batch-1",
        row_number=row_number,
    )


def _outcome(
    lead_id: str = "lead-1",
    session_id: str = "sess-1",
    outcome_status: OutcomeStatus = OutcomeStatus.INTERESTED,
    dnc_status: DNCStatus = DNCStatus.NOT_DNC,
    execution_status: ExecutionStatus = ExecutionStatus.COMPLETED,
    interest_level: InterestLevel = InterestLevel.EXPLICIT_INTEREST,
    problems: tuple[IdentifiedProblem, ...] = (),
    services: tuple[RelevantService, ...] = (),
    objections: tuple[RecordedObjection, ...] = (),
    next_step: str | None = None,
    outstanding_questions: tuple[str, ...] = (),
    summary_lines: tuple[OutcomeSummaryLine, ...] = (),
    termination_reason: TerminationReason = TerminationReason.NATURAL_COMPLETION,
) -> LeadOutcome:
    return LeadOutcome(
        lead_id=lead_id,
        campaign_id="c1",
        session_id=session_id,
        tenant_id="t1",
        outcome_status=outcome_status,
        interest_level=interest_level,
        dnc_status=dnc_status,
        execution_status=execution_status,
        termination_reason=termination_reason,
        identified_problems=problems,
        relevant_services=services,
        objections=objections,
        next_step=next_step,
        outstanding_questions=outstanding_questions,
        summary_lines=summary_lines,
    )


def _write_csv_source(path: str, headers: list[str], rows: list[list[str]]) -> None:
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for row in rows:
            w.writerow(row)


def _write_xlsx_source(path: str, headers: list[str], rows: list[list[str]]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(path)


# ---------------------------------------------------------------------------
# 1. XLSX export
# ---------------------------------------------------------------------------


class TestXLSXExport:
    def test_xlsx_produces_output_file(self, tmp_path) -> None:
        src = str(tmp_path / "leads.xlsx")
        out = str(tmp_path / "leads_results.xlsx")
        _write_xlsx_source(src,
            ["lead_id", "business_name", "phone"],
            [["lead-1", "Alpha Inc", "+15551234567"]],
        )

        result = LeadOutcomeExportWriter().export(
            src, out, (_lead(),), (_outcome(),),
        )
        assert result.export_status == ExportStatus.SUCCESS
        assert os.path.exists(out)
        assert result.output_format == OutputFormat.XLSX

    def test_xlsx_contains_result_columns(self, tmp_path) -> None:
        src = str(tmp_path / "leads.xlsx")
        out = str(tmp_path / "leads_results.xlsx")
        _write_xlsx_source(src,
            ["lead_id", "business_name", "phone"],
            [["lead-1", "Alpha Inc", "+15551234567"]],
        )
        LeadOutcomeExportWriter().export(
            src, out, (_lead(),), (_outcome(),),
        )
        wb = openpyxl.load_workbook(out)
        ws = wb.active
        headers = [c.value for c in ws[1]]
        for col in RESULT_COLUMNS:
            assert col in headers


# ---------------------------------------------------------------------------
# 2. CSV export
# ---------------------------------------------------------------------------


class TestCSVExport:
    def test_csv_produces_output_file(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "leads_results.csv")
        _write_csv_source(src,
            ["lead_id", "business_name", "phone"],
            [["lead-1", "Alpha Inc", "+15551234567"]],
        )
        result = LeadOutcomeExportWriter().export(
            src, out, (_lead(),), (_outcome(),),
        )
        assert result.export_status == ExportStatus.SUCCESS
        assert os.path.exists(out)
        assert result.output_format == OutputFormat.CSV

    def test_csv_result_columns_appended(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "leads_results.csv")
        _write_csv_source(src,
            ["lead_id", "business_name", "phone"],
            [["lead-1", "Alpha Inc", "+15551234567"]],
        )
        LeadOutcomeExportWriter().export(
            src, out, (_lead(),), (_outcome(),),
        )
        with open(out, encoding="utf-8") as f:
            rows = list(csv.reader(f))
        headers = rows[0]
        for col in RESULT_COLUMNS:
            assert col in headers


# ---------------------------------------------------------------------------
# 3. Source file unchanged
# ---------------------------------------------------------------------------


class TestSourceUnchanged:
    def test_xlsx_source_unchanged(self, tmp_path) -> None:
        src = str(tmp_path / "leads.xlsx")
        out = str(tmp_path / "leads_results.xlsx")
        _write_xlsx_source(src,
            ["lead_id", "business_name"],
            [["lead-1", "Alpha Inc"]],
        )
        with open(src, "rb") as f:
            src_bytes_before = f.read()

        LeadOutcomeExportWriter().export(
            src, out, (_lead(),), (_outcome(),),
        )

        with open(src, "rb") as f:
            src_bytes_after = f.read()
        assert src_bytes_before == src_bytes_after

    def test_csv_source_unchanged(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "leads_results.csv")
        _write_csv_source(src,
            ["lead_id", "business_name"],
            [["lead-1", "Alpha Inc"]],
        )
        with open(src, encoding="utf-8") as f:
            src_before = f.read()
        LeadOutcomeExportWriter().export(
            src, out, (_lead(),), (_outcome(),),
        )
        with open(src, encoding="utf-8") as f:
            src_after = f.read()
        assert src_before == src_after


# ---------------------------------------------------------------------------
# 4. Original columns preserved
# ---------------------------------------------------------------------------


class TestOriginalColumnsPreserved:
    def test_csv_original_columns_first(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src,
            ["lead_id", "business_name", "phone", "notes"],
            [["lead-1", "Alpha Inc", "+15551234567", "VIP"]],
        )
        LeadOutcomeExportWriter().export(
            src, out, (_lead(),), (_outcome(),),
        )
        with open(out, encoding="utf-8") as f:
            rows = list(csv.reader(f))
        assert rows[0][:4] == ["lead_id", "business_name", "phone", "notes"]
        assert rows[1][:4] == ["lead-1", "Alpha Inc", "+15551234567", "VIP"]


# ---------------------------------------------------------------------------
# 5. Custom columns preserved
# ---------------------------------------------------------------------------


class TestCustomColumnsPreserved:
    def test_custom_columns_kept(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src,
            ["lead_id", "business_name", "custom_field_xyz", "revenue_bucket"],
            [["lead-1", "Alpha Inc", "foo", "high"]],
        )
        LeadOutcomeExportWriter().export(
            src, out, (_lead(),), (_outcome(),),
        )
        with open(out, encoding="utf-8") as f:
            rows = list(csv.reader(f))
        assert "custom_field_xyz" in rows[0]
        assert "revenue_bucket" in rows[0]
        assert "foo" in rows[1]
        assert "high" in rows[1]


# ---------------------------------------------------------------------------
# 6. Row order preserved
# ---------------------------------------------------------------------------


class TestRowOrderPreserved:
    def test_row_order_preserved(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src,
            ["lead_id", "business_name"],
            [
                ["lead-1", "Alpha"],
                ["lead-2", "Beta"],
                ["lead-3", "Gamma"],
            ],
        )
        leads = (
            _lead("lead-1", row_number=1, business_name="Alpha"),
            _lead("lead-2", row_number=2, business_name="Beta"),
            _lead("lead-3", row_number=3, business_name="Gamma"),
        )
        LeadOutcomeExportWriter().export(src, out, leads, (_outcome("lead-2"),))
        with open(out, encoding="utf-8") as f:
            rows = list(csv.reader(f))
        assert rows[1][1] == "Alpha"
        assert rows[2][1] == "Beta"
        assert rows[3][1] == "Gamma"


# ---------------------------------------------------------------------------
# 7. Correct lead matching
# ---------------------------------------------------------------------------


class TestLeadMatching:
    def test_lead_id_match(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src,
            ["lead_id", "business_name"],
            [["lead-1", "Alpha Inc"]],
        )
        result = LeadOutcomeExportWriter().export(
            src, out, (_lead(),), (_outcome(),),
        )
        assert result.rows_updated == 1

    def test_phone_fallback_match(self, tmp_path) -> None:
        """When no lead_id column exists, phone+campaign is used."""
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src,
            ["business_name", "phone"],
            [["Alpha Inc", "+15551234567"]],
        )
        result = LeadOutcomeExportWriter().export(
            src, out, (_lead(),), (_outcome(),),
        )
        assert result.rows_updated == 1

    def test_row_number_fallback_match(self, tmp_path) -> None:
        """When no lead_id and no phone, row_number is used."""
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src,
            ["business_name"],
            [["Alpha Inc"]],
        )
        result = LeadOutcomeExportWriter().export(
            src, out, (_lead(),), (_outcome(),),
        )
        assert result.rows_updated == 1


# ---------------------------------------------------------------------------
# 8. No fuzzy-name matching
# ---------------------------------------------------------------------------


class TestNoFuzzyMatching:
    def test_similar_names_not_matched(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src,
            ["business_name"],
            [["Alpha Incorporated"]],  # similar name but no row_number/phone/id match
        )
        # Lead has different lead_id and no matching row_number — only business_name is similar
        lead = LeadRecord(
            lead_id="lead-xyz",
            tenant_id="t1",
            campaign_id="c1",
            business_name="Alpha Inc",
            phone=_phone("+19999999999"),
            category=None, city=None, state=None, country=None,
            address=None, email=None, contact_name=None,
            import_batch_id=None,  # no batch
            row_number=None,  # no row match
        )
        result = LeadOutcomeExportWriter().export(
            src, out, (lead,), (_outcome(lead_id="lead-xyz"),),
        )
        # Should NOT match on name similarity
        assert result.rows_updated == 0
        assert result.rows_unmatched == 1


# ---------------------------------------------------------------------------
# 9. Unmatched lead reported
# ---------------------------------------------------------------------------


class TestUnmatchedReporting:
    def test_unmatched_row_in_warnings(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src,
            ["lead_id"],
            [["lead-unknown"]],
        )
        result = LeadOutcomeExportWriter().export(
            src, out, (_lead(),), (_outcome(),),
        )
        assert result.rows_unmatched == 1
        assert any(w.code == "LEAD_NOT_FOUND" for w in result.warnings)


# ---------------------------------------------------------------------------
# 10. Ambiguous match reported
# ---------------------------------------------------------------------------


class TestAmbiguousReporting:
    def test_ambiguous_phone_match(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src,
            ["business_name", "phone"],
            [["Alpha Inc", "+15551234567"]],
        )
        # Two leads share the same phone and no row_number — ambiguous
        lead1 = LeadRecord(
            lead_id="lead-1",
            tenant_id="t1", campaign_id="c1",
            business_name="Alpha", phone=_phone(),
            category=None, city=None, state=None, country=None,
            address=None, email=None, contact_name=None,
            import_batch_id=None, row_number=None,
        )
        lead2 = LeadRecord(
            lead_id="lead-2",
            tenant_id="t1", campaign_id="c1",
            business_name="Alpha2", phone=_phone(),
            category=None, city=None, state=None, country=None,
            address=None, email=None, contact_name=None,
            import_batch_id=None, row_number=None,
        )
        result = LeadOutcomeExportWriter().export(
            src, out, (lead1, lead2),
            (_outcome(lead_id="lead-1"),),
        )
        assert result.ambiguous_matches == 1


# ---------------------------------------------------------------------------
# 11. DNC exported correctly
# ---------------------------------------------------------------------------


class TestDNCExport:
    def test_dnc_outcome_serialized(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src,
            ["lead_id"],
            [["lead-1"]],
        )
        dnc_outcome = _outcome(
            outcome_status=OutcomeStatus.DNC,
            dnc_status=DNCStatus.DNC_CONFIRMED,
            interest_level=InterestLevel.UNKNOWN,
            termination_reason=TerminationReason.DNC_REQUEST,
        )
        LeadOutcomeExportWriter().export(
            src, out, (_lead(),), (dnc_outcome,),
        )
        with open(out, encoding="utf-8") as f:
            rows = list(csv.reader(f))
        headers = rows[0]
        row = rows[1]
        dnc_col = headers.index("DNC Status")
        outcome_col = headers.index("Outcome Status")
        assert row[dnc_col] == "dnc_confirmed"
        assert row[outcome_col] == "dnc"


# ---------------------------------------------------------------------------
# 12. NOT_INTERESTED separate from DNC
# ---------------------------------------------------------------------------


class TestDNCDistinctFromNotInterested:
    def test_not_interested_is_not_dnc(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src,
            ["lead_id"],
            [["lead-1"]],
        )
        ni_outcome = _outcome(
            outcome_status=OutcomeStatus.NOT_INTERESTED,
            dnc_status=DNCStatus.NOT_DNC,
            interest_level=InterestLevel.NO_INTEREST,
            termination_reason=TerminationReason.NOT_INTERESTED_EXIT,
        )
        LeadOutcomeExportWriter().export(
            src, out, (_lead(),), (ni_outcome,),
        )
        with open(out, encoding="utf-8") as f:
            rows = list(csv.reader(f))
        headers = rows[0]
        row = rows[1]
        assert row[headers.index("Outcome Status")] == "not_interested"
        assert row[headers.index("DNC Status")] == "not_dnc"


# ---------------------------------------------------------------------------
# 13. Interested + execution failure preserved
# ---------------------------------------------------------------------------


class TestInterestedPlusFailure:
    def test_interested_execution_failed_both_preserved(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src,
            ["lead_id"],
            [["lead-1"]],
        )
        outcome = _outcome(
            outcome_status=OutcomeStatus.INTERESTED,
            execution_status=ExecutionStatus.FAILED,
            termination_reason=TerminationReason.RUNTIME_ERROR,
        )
        LeadOutcomeExportWriter().export(
            src, out, (_lead(),), (outcome,),
        )
        with open(out, encoding="utf-8") as f:
            rows = list(csv.reader(f))
        headers = rows[0]
        row = rows[1]
        assert row[headers.index("Outcome Status")] == "interested"
        assert row[headers.index("Execution Status")] == "failed"


# ---------------------------------------------------------------------------
# 14. Unknown fields not invented
# ---------------------------------------------------------------------------


class TestUnknownNotInvented:
    def test_unknown_outcome_empty_strings(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src,
            ["lead_id"],
            [["lead-1"]],
        )
        unknown = _outcome(outcome_status=OutcomeStatus.UNKNOWN,
                           interest_level=InterestLevel.UNKNOWN)
        LeadOutcomeExportWriter().export(src, out, (_lead(),), (unknown,))
        with open(out, encoding="utf-8") as f:
            rows = list(csv.reader(f))
        headers = rows[0]
        row = rows[1]
        assert row[headers.index("Next Step")] == ""
        assert row[headers.index("Outstanding Questions")] == ""
        assert row[headers.index("Identified Problems")] == ""


# ---------------------------------------------------------------------------
# 15. Collections serialized deterministically
# ---------------------------------------------------------------------------


class TestCollectionSerialization:
    def test_problems_delimiter_separated(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src, ["lead_id"], [["lead-1"]])
        outcome = _outcome(
            problems=(
                IdentifiedProblem("Zebra problem", FieldConfidence.OBSERVED),
                IdentifiedProblem("Alpha problem", FieldConfidence.OBSERVED),
            ),
            services=(
                RelevantService("svc-z", "match", FieldConfidence.OBSERVED),
                RelevantService("svc-a", "match", FieldConfidence.OBSERVED),
            ),
        )
        LeadOutcomeExportWriter().export(src, out, (_lead(),), (outcome,))
        with open(out, encoding="utf-8") as f:
            rows = list(csv.reader(f))
        headers = rows[0]
        row = rows[1]
        problems_cell = row[headers.index("Identified Problems")]
        services_cell = row[headers.index("Relevant Services")]
        # Deterministically sorted
        assert problems_cell == f"Alpha problem{COLLECTION_DELIMITER}Zebra problem"
        assert services_cell == f"svc-a{COLLECTION_DELIMITER}svc-z"

    def test_no_python_repr_in_cells(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src, ["lead_id"], [["lead-1"]])
        outcome = _outcome(
            objections=(
                RecordedObjection("price", FieldConfidence.OBSERVED),
                RecordedObjection("timing", FieldConfidence.OBSERVED),
            ),
        )
        LeadOutcomeExportWriter().export(src, out, (_lead(),), (outcome,))
        with open(out, encoding="utf-8") as f:
            content = f.read()
        assert "['price'" not in content
        assert "RecordedObjection" not in content


# ---------------------------------------------------------------------------
# 16. Rerun is idempotent
# ---------------------------------------------------------------------------


class TestIdempotency:
    def test_csv_rerun_equivalent(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src, ["lead_id"], [["lead-1"]])
        LeadOutcomeExportWriter().export(src, out, (_lead(),), (_outcome(),))
        with open(out, encoding="utf-8") as f:
            first_run = f.read()
        # Rerun — the output of first run is the input of second? No, source stays the same.
        LeadOutcomeExportWriter().export(src, out, (_lead(),), (_outcome(),))
        with open(out, encoding="utf-8") as f:
            second_run = f.read()
        assert first_run == second_run

    def test_rerun_on_previous_output_as_source_does_not_duplicate_columns(
        self, tmp_path
    ) -> None:
        """If someone accidentally uses the output as input, result columns are not duplicated."""
        src = str(tmp_path / "leads.csv")
        intermediate = str(tmp_path / "intermediate.csv")
        final = str(tmp_path / "final.csv")
        _write_csv_source(src, ["lead_id", "business_name"], [["lead-1", "Alpha"]])
        LeadOutcomeExportWriter().export(src, intermediate, (_lead(),), (_outcome(),))
        # Now export using intermediate as source
        LeadOutcomeExportWriter().export(intermediate, final, (_lead(),), (_outcome(),))
        with open(final, encoding="utf-8") as f:
            rows = list(csv.reader(f))
        headers = rows[0]
        # Result columns should appear exactly once
        for col in RESULT_COLUMNS:
            assert headers.count(col) == 1


# ---------------------------------------------------------------------------
# 17. Result columns not duplicated
# ---------------------------------------------------------------------------


class TestResultColumnsNotDuplicated:
    def test_no_duplicate_result_columns(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src, ["lead_id"], [["lead-1"]])
        LeadOutcomeExportWriter().export(src, out, (_lead(),), (_outcome(),))
        with open(out, encoding="utf-8") as f:
            rows = list(csv.reader(f))
        headers = rows[0]
        for col in RESULT_COLUMNS:
            assert headers.count(col) == 1


# ---------------------------------------------------------------------------
# 18. Multiple outcomes handled safely
# ---------------------------------------------------------------------------


class TestMultipleOutcomes:
    def test_fail_on_ambiguity_by_default(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src, ["lead_id"], [["lead-1"]])
        o1 = _outcome(session_id="sess-1", outcome_status=OutcomeStatus.INTERESTED)
        o2 = _outcome(session_id="sess-2", outcome_status=OutcomeStatus.NOT_INTERESTED)
        result = LeadOutcomeExportWriter().export(
            src, out, (_lead(),), (o1, o2),
        )
        assert result.export_status == ExportStatus.WRITE_FAILED

    def test_latest_trusted_outcome_policy(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src, ["lead_id"], [["lead-1"]])
        o1 = _outcome(session_id="sess-1", outcome_status=OutcomeStatus.INTERESTED)
        o2 = _outcome(session_id="sess-2", outcome_status=OutcomeStatus.NOT_INTERESTED,
                      interest_level=InterestLevel.NO_INTEREST,
                      termination_reason=TerminationReason.NOT_INTERESTED_EXIT)
        result = LeadOutcomeExportWriter().export(
            src, out, (_lead(),), (o1, o2),
            multiple_outcome_policy=MultipleOutcomePolicy.LATEST_TRUSTED_OUTCOME,
            outcome_order_hint=("sess-1", "sess-2"),
        )
        with open(out, encoding="utf-8") as f:
            rows = list(csv.reader(f))
        row = rows[1]
        headers = rows[0]
        assert row[headers.index("Outcome Status")] == "not_interested"

    def test_dnc_wins_over_other_outcomes(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src, ["lead_id"], [["lead-1"]])
        interested = _outcome(session_id="sess-1")
        dnc = _outcome(
            session_id="sess-2",
            outcome_status=OutcomeStatus.DNC,
            dnc_status=DNCStatus.DNC_CONFIRMED,
            interest_level=InterestLevel.UNKNOWN,
            termination_reason=TerminationReason.DNC_REQUEST,
        )
        # DNC always wins — even without latest-trusted policy
        result = LeadOutcomeExportWriter().export(
            src, out, (_lead(),), (interested, dnc),
        )
        # DNC safety invariant: no WRITE_FAILED, DNC preserved
        assert result.export_status in (ExportStatus.SUCCESS, ExportStatus.PARTIAL_SUCCESS)
        with open(out, encoding="utf-8") as f:
            rows = list(csv.reader(f))
        row = rows[1]
        headers = rows[0]
        assert row[headers.index("DNC Status")] == "dnc_confirmed"


# ---------------------------------------------------------------------------
# 19. UTF-8 CSV content
# ---------------------------------------------------------------------------


class TestUTF8CSV:
    def test_utf8_business_name_preserved(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src,
            ["lead_id", "business_name"],
            [["lead-1", "Café Niño"]],
        )
        LeadOutcomeExportWriter().export(src, out, (_lead(),), (_outcome(),))
        with open(out, encoding="utf-8") as f:
            content = f.read()
        assert "Café Niño" in content


# ---------------------------------------------------------------------------
# 20. Workbook formatting preserved reasonably (XLSX)
# ---------------------------------------------------------------------------


class TestWorkbookFormatting:
    def test_header_freeze_pane(self, tmp_path) -> None:
        src = str(tmp_path / "leads.xlsx")
        out = str(tmp_path / "out.xlsx")
        _write_xlsx_source(src,
            ["lead_id", "business_name"],
            [["lead-1", "Alpha"]],
        )
        LeadOutcomeExportWriter().export(src, out, (_lead(),), (_outcome(),))
        wb = openpyxl.load_workbook(out)
        ws = wb.active
        assert ws.freeze_panes == "A2"


# ---------------------------------------------------------------------------
# 21. No diagnostic / internal data leakage
# ---------------------------------------------------------------------------


class TestNoDiagnosticLeakage:
    def test_no_internal_data_in_output(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src, ["lead_id"], [["lead-1"]])
        LeadOutcomeExportWriter().export(src, out, (_lead(),), (_outcome(),))
        with open(out, encoding="utf-8") as f:
            content = f.read()
        for forbidden in (
            "prospect_observed", "prospect_inferred", "business_memory",
            "reasoning", "prompt", "retrieval", "BrainProposal",
            "orchestrator", "ConversationState", "FallbackReason",
        ):
            assert forbidden not in content


# ---------------------------------------------------------------------------
# 22. Deterministic replay
# ---------------------------------------------------------------------------


class TestDeterministicReplay:
    def test_two_runs_produce_byte_identical_csv(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out1 = str(tmp_path / "out1.csv")
        out2 = str(tmp_path / "out2.csv")
        _write_csv_source(src,
            ["lead_id", "business_name"],
            [["lead-1", "Alpha"]],
        )
        outcome = _outcome(
            problems=(
                IdentifiedProblem("P1", FieldConfidence.OBSERVED),
                IdentifiedProblem("P2", FieldConfidence.INFERRED),
            ),
            services=(
                RelevantService("s1", "r1", FieldConfidence.OBSERVED),
                RelevantService("s2", "r2", FieldConfidence.OBSERVED),
            ),
        )
        LeadOutcomeExportWriter().export(src, out1, (_lead(),), (outcome,))
        LeadOutcomeExportWriter().export(src, out2, (_lead(),), (outcome,))
        with open(out1, "rb") as f:
            b1 = f.read()
        with open(out2, "rb") as f:
            b2 = f.read()
        assert b1 == b2


# ---------------------------------------------------------------------------
# Failure modes
# ---------------------------------------------------------------------------


class TestFailureHandling:
    def test_source_not_found(self, tmp_path) -> None:
        result = LeadOutcomeExportWriter().export(
            str(tmp_path / "missing.csv"),
            str(tmp_path / "out.csv"),
            (_lead(),), (_outcome(),),
        )
        assert result.export_status == ExportStatus.SOURCE_FILE_NOT_FOUND

    def test_unsupported_format(self, tmp_path) -> None:
        src = str(tmp_path / "leads.txt")
        with open(src, "w") as f:
            f.write("not a lead file")
        result = LeadOutcomeExportWriter().export(
            src, str(tmp_path / "out.txt"),
            (_lead(),), (_outcome(),),
        )
        assert result.export_status == ExportStatus.UNSUPPORTED_FORMAT


# ---------------------------------------------------------------------------
# Export result contract
# ---------------------------------------------------------------------------


class TestExportResultContract:
    def test_result_counts_zero_on_empty(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src, ["lead_id"], [])
        result = LeadOutcomeExportWriter().export(
            src, out, (), (),
        )
        assert result.rows_total == 0
        assert result.rows_updated == 0

    def test_warning_structure(self, tmp_path) -> None:
        src = str(tmp_path / "leads.csv")
        out = str(tmp_path / "out.csv")
        _write_csv_source(src, ["lead_id"], [["lead-unknown"]])
        result = LeadOutcomeExportWriter().export(
            src, out, (_lead(),), (_outcome(),),
        )
        assert len(result.warnings) >= 1
        w = result.warnings[0]
        assert w.code
        assert w.detail
